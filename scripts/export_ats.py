"""
export_ats.py - Motor de Compilação ATS (Single-Column DOCX & PDF)
Parte do projeto: Career RAG & ATS Intelligence Engine
Autor: Jaques Gabriel Amaral Nunes (https://github.com/jAk7star)
"""

import json
import os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

def load_master_profile(path="data/master_profile.json"):
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)

def export_to_docx(profile, target_role, summary, output_path="Curriculo_ATS_Otimizado.docx"):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    cand = profile["candidate"]
    
    # Header
    h1 = doc.add_paragraph()
    h1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_name = h1.add_run(cand["name"].upper() + "\n")
    run_name.bold = True
    run_name.font.size = Pt(14)
    run_name.font.name = "Arial"
    
    contact_info = f"{cand['location']} | {cand['email']} | LinkedIn: {cand['linkedin']} | GitHub: {cand['github']}"
    run_contact = h1.add_run(contact_info)
    run_contact.font.size = Pt(9.5)
    run_contact.font.name = "Arial"

    def add_section_header(title):
        p = doc.add_paragraph()
        run = p.add_run(title.upper())
        run.bold = True
        run.font.size = Pt(11)
        run.font.name = "Arial"

    add_section_header(f"OBJETIVO: {target_role}")
    
    add_section_header("RESUMO PROFISSIONAL")
    p_sum = doc.add_paragraph(summary)
    p_sum.style.font.name = "Arial"
    p_sum.style.font.size = Pt(10)

    add_section_header("PROJETOS PRÁTICOS E EXPERIÊNCIA TÉCNICA (GITHUB)")
    for proj in profile["projects"]:
        p = doc.add_paragraph(style="List Bullet")
        r_title = p.add_run(f"{proj['name']} ({', '.join(proj['technologies'])}): ")
        r_title.bold = True
        p.add_run(proj["description"])

    add_section_header("FORMAÇÃO ACADÊMICA")
    for edu in profile["education"]:
        doc.add_paragraph(f"{edu['degree']} — {edu['institution']} ({edu['period']})")

    add_section_header("CERTIFICAÇÕES E TREINAMENTOS")
    for cert in profile["certifications"]["tech_and_cloud"]:
        doc.add_paragraph(cert, style="List Bullet")

    doc.save(output_path)
    print(f"[OK] Arquivo DOCX gerado: {output_path}")

def export_to_pdf(profile, target_role, summary, output_path="Curriculo_ATS_Otimizado.pdf"):
    doc = SimpleDocTemplate(output_path, pagesize=A4, leftMargin=40, rightMargin=40, topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()
    story = []

    title_style = ParagraphStyle('NameTitle', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=14, leading=18)
    sub_style = ParagraphStyle('SubInfo', parent=styles['Normal'], fontName='Helvetica', fontSize=9, leading=12)
    sec_style = ParagraphStyle('SecHead', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=11, leading=15, spaceBefore=10, spaceAfter=4)
    body_style = ParagraphStyle('BodyATS', parent=styles['Normal'], fontName='Helvetica', fontSize=9.5, leading=13, spaceAfter=4)

    cand = profile["candidate"]
    story.append(Paragraph(cand["name"].upper(), title_style))
    story.append(Paragraph(f"{cand['location']} | {cand['email']} | {cand['github']}", sub_style))
    story.append(Spacer(1, 8))

    story.append(Paragraph(f"OBJETIVO: {target_role.upper()}", sec_style))
    story.append(Paragraph("RESUMO PROFISSIONAL", sec_style))
    story.append(Paragraph(summary, body_style))

    story.append(Paragraph("PROJETOS PRÁTICOS E PORTFÓLIO (GITHUB)", sec_style))
    for proj in profile["projects"]:
        techs = ", ".join(proj["technologies"])
        story.append(Paragraph(f"• <b>{proj['name']} ({techs}):</b> {proj['description']}", body_style))

    story.append(Paragraph("FORMAÇÃO ACADÊMICA", sec_style))
    for edu in profile["education"]:
        story.append(Paragraph(f"• <b>{edu['degree']}</b> — {edu['institution']} ({edu['period']})", body_style))

    story.append(Paragraph("CERTIFICAÇÕES TECNOLÓGICAS", sec_style))
    for cert in profile["certifications"]["tech_and_cloud"]:
        story.append(Paragraph(f"• {cert}", body_style))

    doc.build(story)
    print(f"[OK] Arquivo PDF gerado: {output_path}")

if __name__ == "__main__":
    profile_data = load_master_profile()
    default_role = "AI Builder Júnior | Engenharia de Prompt & Automação de Dados (Python)"
    default_summary = (
        "Profissional em formação na área de Ciências Exatas (IFAP) com foco prático em construção de soluções de "
        "Inteligência Artificial (LLMs), Engenharia de Prompt e pipelines ETL em Python (Pandas, Regex). "
        "Desenvolvedor da arquitetura Zero-Exposure para sanitização local de dados sensíveis e controle de custos de tokens."
    )
    export_to_docx(profile_data, default_role, default_summary)
    export_to_pdf(profile_data, default_role, default_summary)
